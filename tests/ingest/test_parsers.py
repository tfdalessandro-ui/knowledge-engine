"""Parser tests for every format except PDF (see test_parsers_pdf.py --
PDF parsing needs a network-downloaded OCR/layout model on first use, so it's
gated separately rather than part of the default offline test run)."""
import json

from ingest.parsers import parse_to_text


def test_txt(tmp_path):
    path = tmp_path / "a.txt"
    path.write_text("hello plain text world", encoding="utf-8")
    assert "hello plain text world" in parse_to_text(path)


def test_json(tmp_path):
    path = tmp_path / "a.json"
    path.write_text(json.dumps({"key": "MarkerJSONValue"}), encoding="utf-8")
    assert "MarkerJSONValue" in parse_to_text(path)


def test_xml(tmp_path):
    path = tmp_path / "a.xml"
    path.write_text("<root><item>MarkerXMLValue</item></root>", encoding="utf-8")
    assert "MarkerXMLValue" in parse_to_text(path)


def test_markdown(tmp_path):
    path = tmp_path / "a.md"
    path.write_text("# Title\n\nMarkerMDValue paragraph.", encoding="utf-8")
    assert "MarkerMDValue" in parse_to_text(path)


def test_html(tmp_path):
    path = tmp_path / "a.html"
    path.write_text("<html><body><p>MarkerHTMLValue</p></body></html>", encoding="utf-8")
    assert "MarkerHTMLValue" in parse_to_text(path)


def test_csv(tmp_path):
    path = tmp_path / "a.csv"
    path.write_text("col1,col2\nMarkerCSVValue,2\n", encoding="utf-8")
    assert "MarkerCSVValue" in parse_to_text(path)


def test_docx(tmp_path):
    from docx import Document as DocxDocument

    path = tmp_path / "a.docx"
    doc = DocxDocument()
    doc.add_paragraph("MarkerDOCXValue")
    doc.save(path)
    assert "MarkerDOCXValue" in parse_to_text(path)


def test_xlsx(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "a.xlsx"
    wb = Workbook()
    wb.active["A1"] = "MarkerXLSXValue"
    wb.save(path)
    assert "MarkerXLSXValue" in parse_to_text(path)


def test_pptx(tmp_path):
    from pptx import Presentation

    path = tmp_path / "a.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "MarkerPPTXValue"
    prs.save(path)
    assert "MarkerPPTXValue" in parse_to_text(path)


def test_unsupported_extension_raises(tmp_path):
    path = tmp_path / "a.exe"
    path.write_bytes(b"\x00\x01")
    try:
        parse_to_text(path)
        assert False, "expected ValueError"
    except ValueError:
        pass
